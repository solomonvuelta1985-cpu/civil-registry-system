from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
from copy import deepcopy
import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml import parse_xml


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "documents" / "iScan_CRDMS_Systematic_Presentation_2026_Security_Expanded.pptx"
OUTPUT = ROOT / "documents" / "iScan_CRDMS_Systematic_Presentation_2026_Security_Expanded_v2.pptx"
DEVICE_IMAGE = Path(r"C:\Users\MDRRMO\.codex\generated_images\01a07955-6bfd-7be0-9073-3269c67b3a51\exec-04467c0a-fcce-4a5a-aac0-c01b508a0bb6.png")
REORG_IMAGE = Path(r"C:\Users\MDRRMO\.codex\generated_images\01a07955-6bfd-7be0-9073-3269c67b3a51\exec-3dff5064-2809-4a38-82fa-bad3b9b660e3.png")

NAVY = RGBColor(0x17, 0x26, 0x3A)
MUTED = RGBColor(0x5E, 0x6C, 0x80)
BLUE = RGBColor(0x2F, 0x6B, 0xDE)
TEAL = RGBColor(0x0E, 0x8B, 0x9A)
GREEN = RGBColor(0x16, 0x8A, 0x61)
AMBER = RGBColor(0xB4, 0x53, 0x09)
PALE_BLUE = RGBColor(0xF2, 0xF6, 0xFF)
PALE_AMBER = RGBColor(0xFF, 0xF8, 0xEA)
PALE_GREEN = RGBColor(0xEE, 0xFB, 0xF5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xDA, 0xE2, 0xEE)


def set_run(run, *, size, color, bold=False, font="Arial"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def add_text(slide, x, y, w, h, text="", *, size=16, color=NAVY, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.0):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, color=color, bold=bold)
    return box


def add_rich_lines(slide, x, y, w, h, lines, *, size=15, color=MUTED, gap=6):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        p.level = 0
        for run_text, run_bold, run_color in line:
            run = p.add_run()
            run.text = run_text
            set_run(run, size=size, color=run_color or color, bold=run_bold)
    return box


def add_section_header(slide, kicker, title, subtitle):
    add_text(slide, 0.67, 0.33, 7.3, 0.22, kicker.upper(), size=9.75, color=BLUE, bold=True)
    add_text(slide, 0.67, 0.60, 11.8, 0.52, title, size=27, color=NAVY, bold=True)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.67), Inches(1.31), Inches(0.75), Inches(0.05))
    accent.fill.solid(); accent.fill.fore_color.rgb = BLUE
    accent.line.fill.background()
    add_text(slide, 0.67, 1.58, 11.7, 0.45, subtitle, size=16.5, color=MUTED)


def add_footer(slide, number, total):
    add_text(slide, 0.56, 7.16, 7.3, 0.19, "iScan CRDMS  |  Civil Registry Records Management System", size=8.25, color=RGBColor(0xB8, 0xC9, 0xDE))
    add_text(slide, 11.35, 7.16, 1.42, 0.19, f"{number:02d} / {total:02d}", size=8.25, color=RGBColor(0xB8, 0xC9, 0xDE), bold=True, align=PP_ALIGN.RIGHT)


def add_image_frame(slide, image_path, x, y, w, h):
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    frame.fill.solid(); frame.fill.fore_color.rgb = WHITE
    frame.line.color.rgb = LINE; frame.line.width = Pt(1)
    slide.shapes.add_picture(str(image_path), Inches(x + 0.05), Inches(y + 0.05), width=Inches(w - 0.10), height=Inches(h - 0.10))
    return frame


def add_callout(slide, x, y, w, h, title, body, fill, accent):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = accent; shape.line.width = Pt(1)
    add_text(slide, x + 0.2, y + 0.14, w - 0.4, 0.24, title, size=14, color=accent, bold=True)
    add_text(slide, x + 0.2, y + 0.45, w - 0.4, h - 0.57, body, size=12.5, color=NAVY)
    return shape


def set_notes(slide, text):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    tf.text = text


def make_device_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    add_section_header(
        slide,
        "ADMINISTRATION / ACCESS CONTROL",
        "Registered Devices",
        "Manage which physical devices are allowed to access the system",
    )
    add_image_frame(slide, DEVICE_IMAGE, 7.15, 2.05, 5.55, 3.12)
    add_text(slide, 0.74, 2.08, 5.85, 0.26, "Technical description", size=15.5, color=NAVY, bold=True)
    add_rich_lines(slide, 0.74, 2.43, 5.9, 1.65, [
        [("Device fingerprint registry. ", True, NAVY), ("The system records each registered browser or workstation.", False, MUTED)],
        [("Admin actions. ", True, NAVY), ("Register, review last seen, revoke, or remove a device.", False, MUTED)],
        [("Device Lock. ", True, NAVY), ("When enabled, only approved devices may log in.", False, MUTED)],
    ], size=13.5, gap=6)
    add_callout(
        slide, 0.74, 4.35, 5.9, 1.18,
        "Plain-language explanation",
        "The office keeps a list of allowed computers. An unregistered or revoked device can be blocked when Device Lock is enabled.",
        PALE_GREEN, GREEN,
    )
    add_callout(
        slide, 7.15, 5.38, 5.55, 0.87,
        "Current-state note",
        "The screen shows Device Lock disabled. Register the current device before enabling the restriction.",
        PALE_AMBER, AMBER,
    )
    add_footer(slide, 17, total)
    set_notes(slide, "Registered Devices gives administrators a device registry. The system stores a browser or workstation fingerprint, and an administrator can register, review, revoke, or remove entries. Device Lock is optional, so the screenshot correctly shows that all devices may still log in until an administrator enables the restriction. In simple terms, the office can keep a list of approved computers and block unknown devices after the lock is turned on.")
    return slide


def make_reorg_slide(prs, total):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    add_section_header(
        slide,
        "FILE MAINTENANCE",
        "Reorganize Uploads",
        "Move existing PDFs into the correct year and last-name folder structure",
    )
    add_image_frame(slide, REORG_IMAGE, 0.67, 2.05, 5.65, 3.18)
    add_text(slide, 6.72, 2.08, 5.9, 0.26, "How the tool works", size=15.5, color=NAVY, bold=True)
    add_rich_lines(slide, 6.72, 2.43, 5.8, 1.82, [
        [("Dry Run. ", True, NAVY), ("Scans records and previews the proposed file moves without changing anything.", False, MUTED)],
        [("Apply. ", True, NAVY), ("Moves PDFs and updates database paths after an explicit UNDERSTAND confirmation.", False, MUTED)],
        [("Folder scheme. ", True, NAVY), ("{type}/{year}/{LAST_NAME}/, with the event date taking priority when available.", False, MUTED)],
    ], size=13.5, gap=6)
    add_callout(
        slide, 6.72, 4.58, 5.8, 1.17,
        "Plain-language explanation",
        "It works like a filing-cabinet cleanup: the system shows what it plans to move first, then applies the changes only after confirmation.",
        PALE_BLUE, BLUE,
    )
    add_callout(
        slide, 0.67, 5.52, 5.65, 0.75,
        "Example",
        "birth / 2014 / DELOS_SANTOS / certificate.pdf",
        PALE_AMBER, AMBER,
    )
    add_footer(slide, 18, total)
    set_notes(slide, "Reorganize Uploads helps clean up older PDFs that are in the wrong folder. Dry Run shows the proposed moves without changing files or database paths. Apply performs the move and updates the database only after the user types UNDERSTAND, which gives staff a review step before the change.")
    return slide


def move_live_demo_to_end(prs):
    sld_id_lst = prs.slides._sldIdLst
    if len(sld_id_lst) < 19:
        return
    live_id = sld_id_lst[16]
    sld_id_lst.remove(live_id)
    sld_id_lst.append(live_id)


def update_footer_numbers(prs, total):
    for idx, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame"):
                continue
            for paragraph in shape.text_frame.paragraphs:
                if not paragraph.runs:
                    continue
                whole = "".join(run.text for run in paragraph.runs)
                if re.fullmatch(r"\d{2} / \d{2}", whole.strip()):
                    paragraph.runs[0].text = f"{idx:02d} / {total:02d}"
                    for run in paragraph.runs[1:]:
                        run.text = ""


def add_morph_transitions(path):
    transition = ('<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
                  '<mc:Choice xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" Requires="p159">'
                  '<p:transition spd="slow"><p159:morph option="byObject"/></p:transition>'
                  '</mc:Choice><mc:Fallback><p:transition spd="slow"><p:fade/></p:transition>'
                  '</mc:Fallback></mc:AlternateContent>')
    temp = path.with_suffix('.transition.tmp.pptx')
    with ZipFile(path, 'r') as zin, ZipFile(temp, 'w', ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if re.fullmatch(r'ppt/slides/slide(18|19)\.xml', item.filename):
                text = data.decode('utf-8')
                if '<p:transition' not in text:
                    text = text.replace('</p:sld>', transition + '</p:sld>')
                    data = text.encode('utf-8')
            zout.writestr(item, data)
    temp.replace(path)


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    if not DEVICE_IMAGE.exists() or not REORG_IMAGE.exists():
        raise FileNotFoundError("Generated feature visual missing")

    prs = Presentation(str(SOURCE))
    total = len(prs.slides) + 2
    make_device_slide(prs, total)
    make_reorg_slide(prs, total)
    move_live_demo_to_end(prs)
    update_footer_numbers(prs, total)
    prs.save(str(OUTPUT))
    add_morph_transitions(OUTPUT)
    print(OUTPUT)
    print("slides", len(Presentation(str(OUTPUT)).slides))


if __name__ == "__main__":
    main()
